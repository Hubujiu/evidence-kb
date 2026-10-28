"""Frozen synthetic original-file fixtures for the unified RAG test framework.

No evaluator, model client, budget ledger or experimental controller lives here.
``corpus.json`` contains upload metadata only; ``gold.json`` is scorer-only.
The two splits use separately authored document families and shared distractors
within each split. These are engineering challenges, not human business gold.
"""
from __future__ import annotations

import csv
from datetime import datetime, timezone
from hashlib import sha256
import html
import io
import json
from pathlib import Path
import re
import zipfile

from ..core import PlatformError
from .contracts import AdapterInput, Case, Evidence, Source, Suite, text_hash


VERSION = 'delivery-original-files/1'
FORMATS = ('txt', 'md', 'html', 'csv', 'docx', 'xlsx', 'pptx', 'pdf')


class FixtureParseError(PlatformError):
    def __init__(self, message, *, failures, planned_case_ids):
        super().__init__('FIXTURE_PARSE_EVIDENCE_MISSING', message, 422)
        self.failures = failures
        self.planned_case_ids = list(planned_case_ids)
        failed_documents = {item['document_id'] for item in failures}
        self.failed_document_ids = sorted(failed_documents)


def _stable_zip(raw):
    source, target = io.BytesIO(raw), io.BytesIO()
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target, 'w') as normalized:
        for name in sorted(original.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            payload = original.read(name)
            if name == 'docProps/core.xml':
                payload = re.sub(rb'(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)',
                                 rb'\g<1>2000-01-01T00:00:00Z\g<2>', payload)
            normalized.writestr(info, payload)
    return target.getvalue()


def _file_bytes(format_name, *, lines=(), rows=None):
    if format_name in ('txt', 'md'):
        return ('\n'.join(lines) + '\n').encode('utf-8')
    if format_name == 'html':
        return ('<!doctype html><html><body>' + ''.join('<p>' + html.escape(line) + '</p>' for line in lines) +
                '</body></html>').encode('utf-8')
    if format_name == 'csv':
        stream = io.StringIO(newline='')
        csv.writer(stream, lineterminator='\n').writerows(rows)
        return stream.getvalue().encode('utf-8')
    output = io.BytesIO()
    fixed = datetime(2000, 1, 1, tzinfo=timezone.utc)
    if format_name == 'docx':
        from docx import Document

        document = Document()
        document.core_properties.created = document.core_properties.modified = fixed
        for line in lines:
            document.add_paragraph(line)
        document.save(output)
    elif format_name == 'xlsx':
        from openpyxl import Workbook

        document = Workbook()
        document.properties.created = document.properties.modified = fixed
        document.active.title = '原始记录'
        for row in rows:
            document.active.append(row)
        document.save(output)
    elif format_name == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches

        document = Presentation()
        document.core_properties.created = document.core_properties.modified = fixed
        slide = document.slides.add_slide(document.slide_layouts[6])
        for index, line in enumerate(lines):
            box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5 + index * 0.8), Inches(9), Inches(0.7))
            box.text_frame.text = line
        document.save(output)
    elif format_name == 'pdf':
        from reportlab.pdfgen import canvas

        document = canvas.Canvas(output, pagesize=(680, 400), invariant=1)
        for index, line in enumerate(lines):
            document.drawString(35, 340 - index * 24, line)
        document.showPage()
        document.save()
        return output.getvalue()
    else:
        raise ValueError('Unsupported fixture format')
    return _stable_zip(output.getvalue())


def _specifications():
    """Independent content families; no answer-derived label revisions."""
    return {
        'dev': {
            'family': 'chengchuan-warehouse', 'name': '澄川',
            'txt': '澄川仓储的 E417 表示门禁令牌过期，必须重新登记；未完成身份复核前不得放行。',
            'txt_question': '澄川仓储出现 E417 应怎样处理？放行前还必须满足什么条件？',
            'txt_answer': '重新登记门禁令牌，完成身份复核后才可放行。',
            'reference_title': '澄川停机手册', 'reference_section': '断电检查',
            'reference_rule': '必须切断主电源并确认压力为零，才能拆卸。',
            'html': '冷却箱 C8 运行时禁止开盖；必须断电并冷却10分钟后才能开盖。',
            'html_question': '有人说澄川冷却箱 C8 运行时可直接开盖，这个前提对吗？请给出正确条件。',
            'html_answer': '前提错误。运行时禁止开盖，须断电并冷却10分钟。',
            'orders': [['项目', '物料编号', '数量(件)'], ['澄川', 'A10', '2'], ['澄川', 'A11', '3']],
            'prices': [['项目', '物料编号', '单价(元/件)'], ['澄川', 'A10', '30'], ['澄川', 'A11', '20']],
            'join_answer': '按物料编号连接：A10为2×30=60元，A11为3×20=60元，共120元。',
            'docx': '澄川返修件必须在验收完成后的3个工作日内归档；未验收不得直接归档。',
            'docx_question': '澄川返修件归档的起算条件和期限是什么？未验收能否直接归档？',
            'docx_answer': '验收完成后3个工作日内归档，未验收不得归档。',
            'xlsx': [['记录类型', '保存期限(天)'], ['澄川敏感记录', 180], ['澄川普通记录', 90]],
            'xlsx_question': '澄川敏感记录保存多久？请保留单位并区分普通记录。',
            'xlsx_answer': '敏感记录180天，普通记录90天。',
            'pptx': '澄川值班热线为400-010-7711，服务时段为工作日08:30至18:00。',
            'pptx_question': '澄川值班热线和服务时段分别是什么？',
            'pptx_answer': '400-010-7711，工作日08:30至18:00。',
            'pdf': ['North warehouse entry closes at 17:20.', 'Emergency entry requires supervisor approval.'],
            'pdf_question': 'North warehouse 几点停止正常入仓？紧急入仓是否也要主管批准？',
            'pdf_answer': '17:20停止正常入仓，紧急入仓需要主管批准。',
            'workflow': [['from_step', 'to_step', 'condition_key', 'condition_value'],
                         ['intake', 'review', 'licensed', 'yes'], ['intake', 'return', 'licensed', 'no'],
                         ['review', '__END__', '', '']],
            'workflow_question': '按澄川审核流，当前步骤 intake，已确认 licensed=yes，下一步是什么？',
            'workflow_answer': '下一步为review。',
            'unknown': '澄川公开摘要只记录入库操作，不记录2026年度设备故障率。',
            'unknown_question': '澄川2026年度设备故障率具体是多少？',
            'items': ['营业执照', '经办人证件', '授权委托书', '入库登记表', '货品清单', '质量证明', '运输凭证', '签收单', '安全承诺书'],
            'enumeration_subject': '澄川入库资料清单',
            'old': '澄川加急返修规程2025版：有效期为2025-01-01至2026-09-30，最长处理时限为24小时。',
            'current': '澄川加急返修规程2026版：2026-10-01起生效并替代2025版，最长处理时限为12小时。',
            'draft': '澄川加急返修规程2027草案：尚未批准、不生效，拟将最长处理时限改为8小时。',
            'version_question': '截至2026-10-26，澄川加急返修适用哪一版、时限多久？为什么不能使用2027草案？',
            'version_answer': '适用2026版，最长12小时；2025版已失效，2027草案未批准且不生效。',
        },
        'holdout': {
            'family': 'luanhe-service-onboarding', 'name': '峦河',
            'txt': '峦河服务的 Q209 表示培训凭证尚未同步，需联系教务补录；补录审核通过之前不得授予生产账号。',
            'txt_question': '峦河服务报 Q209 时需要谁做什么？什么时候才能授予生产账号？',
            'txt_answer': '联系教务补录培训凭证，补录审核通过后才能授予生产账号。',
            'reference_title': '峦河交接指引', 'reference_section': '资产核销',
            'reference_rule': '须由保管人与领用人共同清点，并由资产专员签署核销单。',
            'html': '洁净室 R6 清洁后不可立即恢复接待；应完成空气检测且颗粒数达标，连续静置45分钟后方可开放。',
            'html_question': '峦河洁净室 R6 清洁后立刻开放接待是否符合规定？完整恢复条件是什么？',
            'html_answer': '不符合。须空气检测颗粒数达标，并连续静置45分钟后开放。',
            'orders': [['项目', '物料编号', '数量(件)'], ['峦河', 'B21', '4'], ['峦河', 'B24', '2']],
            'prices': [['项目', '物料编号', '单价(元/件)'], ['峦河', 'B21', '15'], ['峦河', 'B24', '25']],
            'join_answer': 'B21为4×15=60元，B24为2×25=50元，共110元。',
            'docx': '峦河外包人员的权限回收应在合同终止后的6小时内完成；事前口头约定不能延长此期限。',
            'docx_question': '峦河外包人员权限回收从何时起算、最晚多久完成？口头约定能延期吗？',
            'docx_answer': '合同终止后6小时内完成，口头约定不能延长。',
            'xlsx': [['记录类型', '保存期限(天)'], ['峦河投诉录音', 365], ['峦河一般回访', 60]],
            'xlsx_question': '峦河投诉录音和一般回访各保存多少天？',
            'xlsx_answer': '投诉录音365天，一般回访60天。',
            'pptx': '峦河工单邮箱为service-desk@luanhe.example，收件时间为周一至周六09:00至20:30。',
            'pptx_question': '峦河工单发到哪个邮箱？周六几点结束收件？',
            'pptx_answer': 'service-desk@luanhe.example，周六20:30结束收件。',
            'pdf': ['South service desk opens at 06:40.', 'Visitor badges must be returned before departure.'],
            'pdf_question': 'South service desk 何时开门？访客离开前对证件有什么要求？',
            'pdf_answer': '06:40开门，离开前必须交回访客证。',
            'workflow': [['from_step', 'to_step', 'condition_key', 'condition_value'],
                         ['handover', 'close', 'assets_clear', 'yes'], ['handover', 'investigate', 'assets_clear', 'no'],
                         ['close', '__END__', '', '']],
            'workflow_question': '按峦河审核流，当前 handover，已确认 assets_clear=no，下一步是什么？',
            'workflow_answer': '下一步为investigate。',
            'unknown': '峦河服务说明仅列受理方式，不提供2026年度净利润数据。',
            'unknown_question': '峦河2026年度净利润具体是多少元？',
            'items': ['资格审查表', '保密协议', '培训合格证', '服务合同', '账号申请单', '设备借用单', '紧急联系人表', '利益冲突申报', '个人信息授权书'],
            'enumeration_subject': '峦河驻场资料清单',
            'old': '峦河客诉升级标准2024版：截至2026-08-31有效，升级等待期限为72小时。',
            'current': '峦河客诉升级标准2026版：自2026-09-01实施，取代2024版，升级等待期限为18小时。',
            'draft': '峦河客诉升级标准2028征求意见稿：未批准实施，仅建议将升级等待期限调整为4小时。',
            'version_question': '2026-10-26峦河客诉升级应等待多久、依据哪版？2028征求意见稿能否作为执行标准？',
            'version_answer': '依据2026版等待18小时；2024版已失效，2028稿未批准实施不能执行。',
        },
    }


def create_delivery_fixtures(output_dir):
    """Write fresh original files and preregister all 24 cases before execution."""
    root = Path(output_dir).resolve()
    root.mkdir(parents=True, exist_ok=False)
    documents, evidence, cases = [], [], []

    def add_document(split, spec, suffix, title, format_name, *, lines=(), rows=None):
        doc_id = f'{split}-{suffix}'
        relative = f'{split}/{title}.{format_name}'
        path = root / relative
        path.parent.mkdir(exist_ok=True)
        payload = _file_bytes(format_name, lines=lines, rows=rows)
        path.write_bytes(payload)
        documents.append({'id': doc_id, 'family': spec['family'], 'split': split, 'title': title,
                          'filename': path.name, 'relative_path': relative, 'format': format_name,
                          'original_sha256': sha256(payload).hexdigest()})
        return doc_id

    def add_evidence(doc_id, quote, suffix='main'):
        eid = doc_id + ':' + suffix
        evidence.append({'id': eid, 'document_id': doc_id, 'quote': quote, 'fact': quote, 'qualifiers': []})
        return eid

    def add_case(split, spec, suffix, question, needed, answer, *, slices=(), behavior='answer'):
        cases.append({'id': split + '-' + suffix, 'family': spec['family'], 'split': split,
            'question': question, 'expected_behavior': behavior, 'any_of': [needed] if needed else [],
            'reference_answer': answer, 'slices': [*slices, 'synthetic_engineering'], 'critical': True})

    def table_text(rows):
        return '\n'.join(' | '.join(str(value) for value in row) for row in rows)
    for split, spec in _specifications().items():
        name = spec['name']
        for format_name in ('txt', 'html', 'docx', 'pptx'):
            doc = add_document(split, spec, format_name, name + '-' + format_name + '-规则', format_name,
                               lines=[spec[format_name]])
            eid = add_evidence(doc, spec[format_name])
            slices = ['format_' + format_name, 'conditions']
            if format_name == 'html':
                slices.append('wrong_premise')
            add_case(split, spec, format_name, spec[format_name + '_question'], [eid],
                     spec[format_name + '_answer'], slices=slices)
        reference = '登记完成后必须遵循《' + spec['reference_title'] + '》§' + spec['reference_section'] + '。'
        seed = add_document(split, spec, 'md', name + '登记说明', 'md', lines=['# ' + name + '登记说明', reference])
        target_text = '§' + spec['reference_section'] + '：' + spec['reference_rule']
        target = add_document(split, spec, 'reference-target', spec['reference_title'], 'md', lines=[target_text])
        add_case(split, spec, 'md', name + '登记完成后所引用的手册规定了哪些操作条件？',
            [add_evidence(seed, reference), add_evidence(target, target_text)], spec['reference_rule'],
            slices=['format_md', 'explicit_reference', 'cross_document'])
        left = add_document(split, spec, 'csv', name + '订单', 'csv', rows=spec['orders'])
        right = add_document(split, spec, 'join-prices', name + '单价', 'csv', rows=spec['prices'])
        add_case(split, spec, 'csv', '请将' + name + '订单与单价按物料编号准确连接，列出各项金额并计算总金额，保留元单位。',
            [add_evidence(left, table_text(spec['orders'])), add_evidence(right, table_text(spec['prices']))],
            spec['join_answer'], slices=['format_csv', 'typed_join', 'numeric', 'units'])
        workbook = add_document(split, spec, 'xlsx', name + '留存标准', 'xlsx', rows=spec['xlsx'])
        add_case(split, spec, 'xlsx', spec['xlsx_question'], [add_evidence(workbook, table_text(spec['xlsx']))],
                 spec['xlsx_answer'], slices=['format_xlsx', 'typed_table', 'units'])
        pdf = add_document(split, spec, 'pdf', name + '英文现场须知', 'pdf', lines=spec['pdf'])
        add_case(split, spec, 'pdf', spec['pdf_question'],
            [add_evidence(pdf, line, str(index)) for index, line in enumerate(spec['pdf'])], spec['pdf_answer'],
            slices=['format_pdf', 'conditions'])
        workflow = add_document(split, spec, 'workflow-source', name + '审核流', 'csv', rows=spec['workflow'])
        chosen_edge = spec['workflow'][1] if split == 'dev' else spec['workflow'][2]
        workflow_evidence = [add_evidence(workflow, table_text([spec['workflow'][0]]), 'headers'),
                             add_evidence(workflow, table_text([chosen_edge]), 'selected-edge')]
        add_case(split, spec, 'workflow', spec['workflow_question'], workflow_evidence,
                 spec['workflow_answer'], slices=['workflow', 'typed_table', 'branch_state'])
        unknown = add_document(split, spec, 'unknown-source', name + '公开摘要', 'txt', lines=[spec['unknown']])
        add_evidence(unknown, spec['unknown'])
        add_case(split, spec, 'no-answer', spec['unknown_question'], [], '现有资料未提供该数值，应明确说明证据不足。',
                 slices=['no_answer'], behavior='abstain')
        item_evidence = []
        for index, item in enumerate(spec['items'], 1):
            statement = spec['enumeration_subject'] + f'第{index}项：' + item + '。此条属于完整九项清单。'
            document = add_document(split, spec, 'all-item-' + str(index), spec['enumeration_subject'] + f'-{index}',
                                    'txt', lines=[statement])
            item_evidence.append(add_evidence(document, statement))
        add_case(split, spec, 'enumeration', '请完整列出' + spec['enumeration_subject'] + '的全部九项，不得只给前几项。',
                 item_evidence, '、'.join(spec['items']), slices=['exhaustive', 'more_sources_than_top5'])
        version_evidence = []
        for version in ('old', 'current', 'draft'):
            document = add_document(split, spec, 'version-' + version, name + '版本标准-' + version, 'md',
                                    lines=[spec[version]])
            version_evidence.append(add_evidence(document, spec[version]))
        add_case(split, spec, 'validity', spec['version_question'], version_evidence[1:], spec['version_answer'],
                 slices=['business_validity', 'superseded', 'draft'])
        for index in range(6):
            statement = f'远港试验项目{index}的留存期限为{20 + index}天；只适用远港，不适用{name}。'
            add_document(split, spec, 'distractor-' + str(index), f'远港-{split}-试验记录-{index}', 'txt',
                         lines=[statement, '一般说明：请联系项目管理员。', '一般说明：请联系项目管理员。'])
    manifest = {'schema_version': VERSION, 'label_provenance': 'synthetic_engineering',
        'documents': documents, 'evidence': evidence, 'cases': cases,
        'scope_policy': 'one_shared_mixed_corpus_per_split; all_split_documents_allowed_for_every_case',
        'quality_claim': 'synthetic_engineering_challenge_only; not_human_enterprise_gold'}
    scenarios = {'dev-workflow': {'workflow_state': {'current_step': 'intake', 'facts': {'licensed': 'yes'}}},
                 'holdout-workflow': {'workflow_state': {'current_step': 'handover', 'facts': {'assets_clear': 'no'}}}}
    for filename, data in (('corpus.json', {'schema_version': VERSION, 'documents': documents}),
                           ('scenario-inputs.json', scenarios),
                           ('gold.json', {'schema_version': VERSION, 'evidence': evidence, 'cases': cases}),
                           ('fixtures.json', manifest)):
        (root / filename).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    return {**manifest, 'root': str(root)}


def load_delivery_fixtures(root):
    root = Path(root).resolve()
    manifest = json.loads((root / 'fixtures.json').read_text('utf-8'))
    if manifest.get('schema_version') != VERSION:
        raise ValueError('Unknown fixture version')
    for document in manifest['documents']:
        path = (root / document['relative_path']).resolve()
        if not path.is_relative_to(root) or sha256(path.read_bytes()).hexdigest() != document['original_sha256']:
            raise ValueError('Original fixture hash mismatch: ' + document['id'])
    return {**manifest, 'root': str(root)}


def bind_suite(platform, actor, scope_id, fixtures, ingested, *, split=None, context_budget=16000,
               provider_id='magpie_luna'):
    """Bind preregistered gold to exact immutable parses after real upload/build.

    ``ingested`` maps fixture document IDs to actual Platform ingest receipts.
    The returned citation directory contains original text and offsets only,
    and is safe for an adapter. ``suite`` and ``expectations`` are scorer-only.
    Missing/ambiguous preregistered quotations are parse failures; labels and
    the fixed case denominator are never edited to accommodate parser output.
    """
    from ..agent_tools import EvidenceGateway
    from ..retrieval_pipeline import _allowed_bindings, _catalog

    if split not in (None, 'dev', 'holdout') or type(context_budget) is not int or context_budget < 1:
        raise ValueError('Invalid fixture split or context budget')
    fixtures = load_delivery_fixtures(fixtures['root'] if isinstance(fixtures, dict) else fixtures)
    documents = [doc for doc in fixtures['documents'] if split is None or doc['split'] == split]
    cases = [case for case in fixtures['cases'] if split is None or case['split'] == split]
    planned_case_ids = [case['id'] for case in cases]
    gateway = EvidenceGateway(platform, provider_id=provider_id)
    state, bindings, _ = gateway._state(actor, scope_id, with_build=True)
    allowed = _allowed_bindings(gateway, actor, bindings, {'mode': 'AUTO', 'source_ids': [],
        'versions': {}, 'version_policy': 'latest', 'lifetime': 'session'})
    catalog, units = _catalog(gateway, actor, scope_id, state, allowed)
    by_source = {doc['source_id']: doc for doc in catalog}
    actual, sources, failures = {}, [], []
    for document in documents:
        receipt = ingested.get(document['id'])
        source = by_source.get(receipt.get('source_id')) if isinstance(receipt, dict) else None
        if (source is None or source['source_version_id'] != document['original_sha256'] or
                source['document'].get('original_sha256') != document['original_sha256'] or
                source['document'].get('status') != 'READY'):
            failures.append({'document_id': document['id'], 'reason': 'missing_wrong_version_or_incomplete_parse'})
            continue
        actual[document['id']] = source
        sources.append(Source(id=source['source_id'], kb=scope_id, version=source['source_version_id'],
                              text=source['text'], sha256=text_hash(source['text'])))
    bound_evidence = []
    selected_doc_ids = {doc['id'] for doc in documents}
    for annotation in fixtures['evidence']:
        if annotation['document_id'] not in selected_doc_ids or annotation['document_id'] not in actual:
            continue
        source = actual[annotation['document_id']]
        quote = annotation['quote']
        # Annotation binding is scorer-only. Adapter citation offsets below
        # come from immutable element positions, never substring matching.
        if source['text'].count(quote) != 1:
            failures.append({'document_id': annotation['document_id'], 'evidence_id': annotation['id'],
                             'reason': 'preregistered_quote_missing_or_ambiguous', 'expected_quote': quote})
            continue
        start = source['text'].index(quote)
        bound_evidence.append(Evidence(id=annotation['id'], source_id=source['source_id'], start=start,
            end=start + len(quote), quote=quote, fact=annotation['fact'], qualifiers=tuple(annotation['qualifiers'])))
    if failures:
        raise FixtureParseError('Prerecorded original evidence is missing; retain planned cases as parse failures',
                                failures=failures, planned_case_ids=planned_case_ids)
    source_ids = tuple(source.id for source in sources)
    if set(by_source) != set(source_ids):
        raise PlatformError('FIXTURE_SCOPE_MISMATCH',
                            'Use one complete shared fixture corpus per scope; do not mix untracked sources or splits', 409)
    bound_cases = []
    evidence_by_id = {item.id: item for item in bound_evidence}
    for case in cases:
        relevant = {evidence_by_id[eid].source_id for group in case['any_of'] for eid in group}
        bound_cases.append(Case(input=AdapterInput(case_id=case['id'], question=case['question'], kb=scope_id,
            principal=actor.principal_id, acl_revision=str(state['authority_revision']), allowed_sources=source_ids,
            as_of='2026-10-26', context_budget=context_budget), family=case['family'], split=case['split'],
            slices=tuple(case['slices']), expected_behavior=case['expected_behavior'],
            any_of=tuple(tuple(group) for group in case['any_of']), qrels={source: 1 for source in sorted(relevant)}))
    suite = Suite(id=VERSION + '-' + (split or 'all'), label_provenance='synthetic_engineering',
                  sources=tuple(sources), evidence=tuple(bound_evidence), cases=tuple(bound_cases))
    directory = {}
    selected_sources = set(source_ids)
    for citation_id, unit in units.items():
        source = unit['source']
        if source['source_id'] not in selected_sources:
            continue
        canonical = gateway.read(actor, scope_id, [citation_id], {})[0]
        if canonical['text'] != source['text'][unit['start']:unit['end']]:
            raise PlatformError('FIXTURE_CITATION_MAPPING_MISMATCH', 'Canonical citation differs from immutable parse span', 503)
        directory[citation_id] = {**canonical, 'kind': unit['element']['kind'],
            'parse_text_start': unit['start'], 'parse_text_end': unit['end'],
            'source_text_sha256': text_hash(source['text'])}
    gateway._unchanged(actor, scope_id, state)
    return {'suite': suite, 'bindings': directory,
        'documents': {doc['id']: {'source_id': actual[doc['id']]['source_id'],
            'source_version_id': actual[doc['id']]['source_version_id'],
            'parse_snapshot_id': actual[doc['id']]['parse_snapshot_id'],
            'original_sha256': doc['original_sha256']} for doc in documents},
        'expectations': {case['id']: {'reference_answer': case['reference_answer'],
            'expected_behavior': case['expected_behavior'], 'slices': case['slices']} for case in cases},
        'planned_case_ids': planned_case_ids, 'manifest_id': state['manifest_id'], 'parse_failures': []}
