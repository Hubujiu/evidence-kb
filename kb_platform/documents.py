"""Native document parsing with immutable, explicitly graded source anchors."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import zipfile
from datetime import date, datetime
from importlib.metadata import version
from pathlib import PurePath
from typing import Any

from .markdown_source import PARSER_VERSION as TEXT_PARSER_VERSION, source_blocks


PARSER_VERSION = "native-evidence/0.2.1"
MAX_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED = 100 * 1024 * 1024
MAX_PAGES = 300
MAX_ELEMENTS = 30_000
MAX_CELLS = 100_000
MEDIA_TYPES = {
    ".txt": "text/plain", ".md": "text/markdown", ".csv": "text/csv",
    ".html": "text/html", ".htm": "text/html", ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".tif": "image/tiff",
    ".tiff": "image/tiff",
}
DEPENDENCIES = {
    ".md": ("markdown-it-py",), ".txt": ("markdown-it-py",),
    ".pdf": ("pdfplumber", "pdfminer.six", "pypdfium2"),
    ".docx": ("python-docx",), ".pptx": ("python-pptx",), ".xlsx": ("openpyxl",),
    ".html": ("beautifulsoup4",), ".htm": ("beautifulsoup4",),
    ".png": ("Pillow",), ".jpg": ("Pillow",), ".jpeg": ("Pillow",),
    ".tif": ("Pillow",), ".tiff": ("Pillow",),
}


class DocumentError(ValueError):
    pass


def max_upload_bytes() -> int:
    """One explicit byte limit for HTTP admission and the parser subprocess."""
    raw = os.environ.get("KB_MAX_UPLOAD_BYTES", str(MAX_BYTES))
    try:
        limit = int(raw)
    except (TypeError, ValueError):
        raise ValueError("KB_MAX_UPLOAD_BYTES must be a positive integer") from None
    if limit <= 0:
        raise ValueError("KB_MAX_UPLOAD_BYTES must be a positive integer")
    return limit


def _hash(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def _normalize(raw: str) -> tuple[str, list[int]]:
    chars, positions = [], []
    index = 0
    while index < len(raw):
        positions.append(index)
        if raw[index:index + 2] == "\r\n":
            chars.append("\n")
            index += 2
        else:
            chars.append("\n" if raw[index] == "\r" else raw[index])
            index += 1
    positions.append(len(raw))
    return "".join(chars), positions


def _json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _text_kind(text: str) -> str:
    stripped = text.strip()
    if re.match(r"^(?:警告|注意|WARNING|CAUTION)\s*[:：]", stripped, re.I):
        return "Warning"
    if re.match(r"^(?:前提|先决条件|Prerequisite)\s*[:：]", stripped, re.I):
        return "Prerequisite"
    if re.match(r"^\d+[.)、]\s*\S", stripped):
        return "OrderedStep"
    return "Paragraph"


class _Bundle:
    def __init__(self, data: bytes, filename: str, policy: dict):
        self.policy = policy
        self.ext = PurePath(filename).suffix.lower()
        self.source = hashlib.sha256(data).hexdigest()
        packages = {name: version(name) for name in DEPENDENCIES.get(self.ext, ())}
        config = {key: policy.get(key, default) for key, default in {
            "native_parse": True, "ocr": False, "include_hidden": False, "include_notes": False,
        }.items()}
        parser_version = TEXT_PARSER_VERSION if self.ext in (".md", ".txt") else PARSER_VERSION
        self.snapshot = _hash([self.source, parser_version, packages, config, self.ext])
        self.result = {
            "source_version_id": self.source, "original_sha256": self.source,
            "filename": PurePath(filename).name, "media_type": MEDIA_TYPES.get(self.ext, "application/octet-stream"),
            "parse_snapshot_id": self.snapshot, "parser_version": parser_version, "status": "READY",
            "elements": [], "evidence_spans": [], "anchors": [],
            "quality_report": {
                "warnings": [], "missing_capabilities": [], "excluded": [], "ocr_decisions": [],
                "dependency_versions": packages, "normalization": "CRLF/CR to LF only",
                "offset_unit": "unicode_code_point", "model_calls": 0,
            },
        }
        self.headings: list[tuple[str, str]] = []

    def identifier(self, kind: str, location: Any) -> str:
        return f"{kind}_{_hash([self.snapshot, kind, location])[:32]}"

    def warning(self, code: str, **details):
        self.result["quality_report"]["warnings"].append({"code": code, **details})

    def missing(self, capability: str, **details):
        self.result["status"] = "PARTIAL"
        missing = self.result["quality_report"]["missing_capabilities"]
        if capability not in missing:
            missing.append(capability)
        self.warning(capability, **details)

    def anchor(self, text: str, location: dict) -> dict:
        result = {
            "anchor_id": self.identifier("anchor", location), "source_version_id": self.source,
            "quote": text, "offset_unit": "unicode_code_point", "text_start": 0, "text_end": len(text),
            "granularity": "object", "coordinate_quality": "object_exact", **location,
        }
        self.result["anchors"].append(result)
        return result

    def element(self, text: str, location: dict, kind: str | None = None, **extra) -> dict:
        if len(self.result["elements"]) >= MAX_ELEMENTS:
            raise DocumentError("element_limit")
        normalized, mapping = _normalize(text)
        anchor = self.anchor(text, location)
        element_id = self.identifier("element", location)
        evidence_id = self.identifier("evidence", location)
        element = {
            "element_id": element_id, "kind": kind or _text_kind(text),
            "ordinal": len(self.result["elements"]), "raw_text": text, "normalized_text": normalized,
            "normalization_map": mapping, "heading_path": [item[0] for item in self.headings],
            "parent_id": self.headings[-1][1] if self.headings else None,
            "evidence_ids": [evidence_id], "extraction_method": "native", **extra,
        }
        self.result["elements"].append(element)
        self.result["evidence_spans"].append({
            "evidence_id": evidence_id, "source_version_id": self.source, "parse_snapshot_id": self.snapshot,
            "element_ids": [element_id], "anchor_ids": [anchor["anchor_id"]],
            "extraction_method": "native", "native_text_spans": [[0, len(text)]],
            "normalized_text_spans": [[0, len(normalized)]],
        })
        return element

    def heading(self, text: str, level: int, location: dict):
        self.headings = self.headings[:max(0, level - 1)]
        element = self.element(text, location, "Section", heading_level=level)
        self.headings.append((text, element["element_id"]))

    def table(self, grid: list[list[Any]], location: dict, cell_locations: list[list[dict]],
              cell_details: dict[tuple[int, int], dict] | None = None, merged_ranges=None):
        if not grid:
            return
        if sum(map(len, grid)) > MAX_CELLS:
            raise DocumentError("cell_limit")
        cells, cell_anchors = [], []
        details = cell_details or {}
        for row_index, row in enumerate(grid):
            for col_index, value in enumerate(row):
                raw = "" if value is None else str(value)
                metadata = details.get((row_index, col_index), {})
                anchor = self.anchor(raw, cell_locations[row_index][col_index])
                cell_anchors.append(anchor["anchor_id"])
                cells.append({
                    "row": row_index + 1, "column": col_index + 1, "raw_text": raw,
                    "parsed_value": _json_value(value), "formula": None, "cached_value": None,
                    "value_status": "native", **metadata, "cell_anchor": anchor["anchor_id"],
                })
        text = "\n".join(" | ".join("" if v is None else str(v) for v in row) for row in grid)
        element = self.element(text, location, "Table", table={
            "headers": ["" if v is None else str(v) for v in grid[0]],
            "rows": [[_json_value(v) for v in row] for row in grid[1:]],
            "cells": cells, "merged_ranges": merged_ranges or [],
        })
        self.result["evidence_spans"][-1]["anchor_ids"].extend(cell_anchors)
        return element

    def ocr_needed(self, location: dict, native_text: bool):
        decision = "OCR_REGION" if native_text else "OCR_PAGE"
        if not self.policy.get("ocr", False):
            decision = "BLOCKED_BY_POLICY"
        self.result["quality_report"]["ocr_decisions"].append({"decision": decision, **location})
        self.missing("ocr_not_executed", **location)


def _validate_zip(data: bytes, ext: str):
    required = {".docx": "word/document.xml", ".pptx": "ppt/presentation.xml", ".xlsx": "xl/workbook.xml"}[ext]
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items = archive.infolist()
        if len(items) > 10_000 or sum(item.file_size for item in items) > MAX_UNCOMPRESSED:
            raise DocumentError("archive_size_limit")
        for item in items:
            if item.file_size > max(1024 * 1024, item.compress_size * 200):
                raise DocumentError("archive_ratio_limit")
            if item.flag_bits & 1:
                raise DocumentError("encrypted_archive")
            if "vbaproject" in item.filename.lower() or "/embeddings/" in item.filename.lower():
                raise DocumentError("active_or_embedded_content")
        if required not in archive.namelist() or "[Content_Types].xml" not in archive.namelist():
            raise DocumentError("file_type_mismatch")


def _decode(data: bytes) -> str:
    text = data.decode("utf-8-sig")
    if "\x00" in text:
        raise DocumentError("binary_text_input")
    return text


def _parse_text(bundle: _Bundle, data: bytes):
    text = _decode(data)
    for block in source_blocks(text, markdown=bundle.ext == ".md"):
        content = text[block.start:block.end]
        path = f"line/{block.first_line}"
        if block.unclosed and block.end == len(text):
            path += "-end"
        elif block.last_line != block.first_line or block.kind == "Code":
            path += f"-{block.last_line}"
        location = {
            "selector_type": "text_position", "object_path": path,
            "text_start": block.start, "text_end": block.end,
            "granularity": "block" if "-" in path else "line",
        }
        if block.heading_level:
            bundle.heading(content, block.heading_level, location)
        else:
            bundle.element(content, location, block.kind)
        if block.unclosed:
            bundle.warning("unclosed_code_fence")


def _parse_csv(bundle: _Bundle, data: bytes):
    grid = list(csv.reader(io.StringIO(_decode(data), newline="")))
    width = max(map(len, grid), default=0)
    grid = [row + [None] * (width - len(row)) for row in grid]
    locations = [[{"selector_type": "cell", "sheet": "CSV", "cell_range": f"R{r+1}C{c+1}",
                   "object_path": f"row/{r+1}/cell/{c+1}", "granularity": "cell"}
                  for c in range(width)] for r in range(len(grid))]
    bundle.table(grid, {"selector_type": "cell_range", "sheet": "CSV", "cell_range": f"R1C1:R{len(grid)}C{width}",
                        "object_path": "table/1", "granularity": "table"}, locations)


def _parse_html(bundle: _Bundle, data: bytes):
    from bs4 import BeautifulSoup, Comment, NavigableString, Tag

    soup = BeautifulSoup(_decode(data), "html.parser")
    paths = {}
    for tag in soup.find_all(True):
        ancestors = [tag, *list(tag.parents)]
        components = []
        for ancestor in reversed(ancestors):
            if ancestor.name == "[document]":
                continue
            position = 1 + sum(1 for sibling in ancestor.previous_siblings
                               if isinstance(sibling, Tag) and sibling.name == ancestor.name)
            components.append(f"{ancestor.name}[{position}]")
        paths[id(tag)] = "/" + "/".join(components)
    for dangerous in soup(["script", "style", "noscript", "iframe", "object", "embed", "template"]):
        dangerous.decompose()
    if not bundle.policy.get("include_hidden", False):
        for hidden in list(soup.find_all(True)):
            if hidden.attrs is not None and (hidden.has_attr("hidden") or
                re.search(r"(?:^|;)\s*(?:display\s*:\s*none|visibility\s*:\s*hidden)\s*(?:!important\s*)?(?:;|$)",
                          hidden.get("style", ""), re.I)):
                hidden.decompose()

    block_names = {"html", "body", "div", "section", "article", "aside", "header", "footer", "main",
                   "blockquote", "p", "li", "pre", "table", "ul", "ol", "dl", "dt", "dd", "address",
                   "figure", "figcaption", "details", "summary", "h1", "h2", "h3", "h4", "h5", "h6"}

    def location(tag, **extra):
        return {"selector_type": "object_path", "object_path": paths.get(id(tag), "/"), **extra}

    def visible_text(tag):
        return "".join("\n" if isinstance(node, Tag) and node.name == "br" else str(node)
                       for node in tag.descendants
                       if (isinstance(node, NavigableString) and not isinstance(node, Comment)) or
                       (isinstance(node, Tag) and node.name == "br"))

    def visit(tag):
        if tag.name == "table":
            caption = tag.find("caption", recursive=False)
            if caption is not None and visible_text(caption).strip():
                bundle.element(visible_text(caption), location(caption), "Caption")
            cells = [row.find_all(["td", "th"], recursive=False) for row in tag.find_all("tr")
                     if row.find_parent("table") is tag]
            grid = [[visible_text(cell) for cell in row] for row in cells]
            locations = [[location(cell, granularity="cell") for cell in row] for row in cells]
            bundle.table(grid, location(tag, granularity="table"), locations)
            if tag.find(attrs={"rowspan": True}) or tag.find(attrs={"colspan": True}):
                bundle.missing("html_merged_grid_requires_review")
            if tag.find("table"):
                bundle.missing("html_nested_table_requires_review")
        elif re.fullmatch(r"h[1-6]", tag.name):
            bundle.heading(visible_text(tag), int(tag.name[1]), location(tag))
        elif tag.name == "pre":
            bundle.element(visible_text(tag), location(tag), "Code")
        else:
            fragments, run_index = [], 0

            def flush():
                nonlocal run_index
                text = "".join(fragments)
                fragments.clear()
                if text.strip():
                    bundle.element(text, location(tag, text_run_index=run_index))
                    run_index += 1

            for child in tag.children:
                if isinstance(child, Comment):
                    continue
                if isinstance(child, NavigableString):
                    fragments.append(str(child))
                elif child.name == "br":
                    fragments.append("\n")
                elif child.name in block_names or child.find(block_names):
                    flush()
                    visit(child)
                elif child.name not in {"head", "img"}:
                    fragments.append(visible_text(child))
            flush()

    visit(soup.body or soup)
    if soup.find("img"):
        bundle.missing("html_images_not_fetched")


def _pdf_location(page, index: int, bbox: tuple | list, suffix: str) -> dict:
    return {
        "selector_type": "page_region", "page_index": index, "page_label": str(index + 1),
        "rects": [list(map(float, bbox))], "bbox": list(map(float, bbox)),
        "coordinate_system": "pdfplumber_top_left_points", "dimensions": [float(page.width), float(page.height)],
        "page_bbox": list(map(float, page.bbox)), "crop_box": list(map(float, page.cropbox or page.bbox)),
        "rotation": page.rotation, "object_path": f"page/{index}/{suffix}",
        "granularity": "line", "coordinate_quality": "native_geometry",
    }


def _parse_pdf(bundle: _Bundle, data: bytes):
    import pdfplumber

    if not data.startswith(b"%PDF-"):
        raise DocumentError("file_type_mismatch")
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        if len(pdf.pages) > MAX_PAGES:
            raise DocumentError("page_limit")
        for index, page in enumerate(pdf.pages):
            rotated_page = (
                page.rotation in (90, 270) and page.chars and all(not char["upright"] for char in page.chars)
            ) or (page.rotation == 180 and page.chars and all(char["matrix"][0] < 0 for char in page.chars))
            if rotated_page:
                from pdfplumber.utils import cluster_objects

                direction = "ttb" if page.rotation == 90 else "btt"
                if page.rotation == 180:
                    words = page.extract_words(line_dir="btt", char_dir="rtl")
                    groups = cluster_objects(words, "top", 3)
                else:
                    words = page.extract_words(line_dir_rotated="ltr", char_dir_rotated=direction)
                    groups = cluster_objects(words, "x0", 3)
                if page.rotation in (90, 180):
                    groups.reverse()
                lines = []
                for group in groups:
                    sort_key = "x0" if page.rotation == 180 else "top"
                    ordered = sorted(group, key=lambda word: word[sort_key], reverse=direction == "btt")
                    lines.append({"text": " ".join(word["text"] for word in ordered),
                                  "x0": min(word["x0"] for word in group), "x1": max(word["x1"] for word in group),
                                  "top": min(word["top"] for word in group), "bottom": max(word["bottom"] for word in group)})
            else:
                lines = page.extract_text_lines(return_chars=False)
            for line_index, line in enumerate(lines):
                bbox = (line["x0"], line["top"], line["x1"], line["bottom"])
                bundle.element(line["text"], _pdf_location(page, index, bbox, f"line/{line_index}"))
            for table_index, table in enumerate(page.find_tables()):
                grid = table.extract()
                cell_locations = []
                for r, row in enumerate(table.rows):
                    cell_locations.append([{
                        **_pdf_location(page, index, cell or table.bbox, f"table/{table_index}/cell/{r}/{c}"),
                        "granularity": "cell" if cell else "table", "coordinate_quality": "native_geometry" if cell else "coarse",
                    } for c, cell in enumerate(row.cells)])
                bundle.table(grid, {**_pdf_location(page, index, table.bbox, f"table/{table_index}"),
                                    "granularity": "table"}, cell_locations)
            if page.images:
                bundle.ocr_needed({"page_index": index, "image_regions": len(page.images)}, bool(lines))
            elif not lines and page.curves:
                bundle.result["quality_report"]["ocr_decisions"].append({"page_index": index, "decision": "REVIEW_REQUIRED"})
                bundle.missing("vector_only_page_requires_review", page_index=index)
            else:
                bundle.result["quality_report"]["ocr_decisions"].append({"page_index": index, "decision": "SKIP_NATIVE_OK"})
            if any("\ufffd" in line["text"] or "(cid:" in line["text"] for line in lines):
                bundle.missing("native_text_quality_requires_review", page_index=index)
            page.close()


def _docx_text(paragraph, include_hidden: bool) -> str:
    if include_hidden:
        return paragraph.text

    def hidden(run):
        if run.font.hidden is not None:
            return run.font.hidden
        for initial in (run.style, paragraph.style):
            style, visited = initial, set()
            while style is not None and style.style_id not in visited:
                visited.add(style.style_id)
                if style.font.hidden is not None:
                    return style.font.hidden
                style = style.base_style
        return False

    runs = (run for item in paragraph.iter_inner_content()
            for run in (item.runs if hasattr(item, "runs") else [item]))
    return "".join(run.text for run in runs if not hidden(run))


def _parse_docx(bundle: _Bundle, data: bytes):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    include_hidden = bundle.policy.get("include_hidden", False)
    nested_cells = 0

    def cell_text(cell, object_path: str, depth: int = 0) -> str:
        """Preserve nested tables in XML order, anchored to their containing cell.

        Nested grids remain text inside the parent table cell; they are not
        advertised as independent table-query objects or page coordinates.
        """
        nonlocal nested_cells
        if depth > 32:
            raise DocumentError("nested_table_depth_limit")
        parts = []
        for inner_index, item in enumerate(cell.iter_inner_content()):
            if isinstance(item, Paragraph):
                parts.append(_docx_text(item, include_hidden))
            elif isinstance(item, Table):
                nested_path = f"{object_path}/table/{inner_index}"
                nested_grid = []
                for row_index, nested_row in enumerate(item.rows):
                    nested_cells += len(nested_row.cells)
                    if nested_cells > MAX_CELLS:
                        raise DocumentError("cell_limit")
                    nested_grid.append(" | ".join(
                        cell_text(nested_cell, f"{nested_path}/row/{row_index}/cell/{column_index}", depth + 1)
                        for column_index, nested_cell in enumerate(nested_row.cells)))
                parts.append("\n".join(nested_grid))
                bundle.warning("nested_docx_table_flattened", object_path=nested_path,
                               anchor_granularity="containing_cell", independent_table_query=False)
        return "\n".join(parts)

    for index, child in enumerate(doc.element.body):
        location = {"selector_type": "object_path", "object_path": f"/word/document.xml/body/{index}"}
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, doc)
            text = _docx_text(paragraph, include_hidden)
            if not text.strip():
                continue
            heading = re.match(r"Heading ([1-6])$", paragraph.style.name)
            if heading:
                bundle.heading(text, int(heading[1]), location)
            else:
                bundle.element(text, location)
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            grid, locations, merged, seen = [], [], [], {}
            for r, row in enumerate(table.rows):
                grid.append([])
                locations.append([])
                for c, cell in enumerate(row.cells):
                    cell_path = f"{location['object_path']}/row/{r}/cell/{c}"
                    grid[-1].append(cell_text(cell, cell_path))
                    locations[-1].append({"selector_type": "object_path", "object_path": cell_path, "granularity": "cell"})
                    if cell._tc in seen:
                        merged.append({"cell": [r+1, c+1], "origin": seen[cell._tc]})
                    else:
                        seen[cell._tc] = [r+1, c+1]
            bundle.table(grid, {**location, "granularity": "table"}, locations, merged_ranges=merged)
    if doc.inline_shapes:
        bundle.missing("office_images_not_recognized", count=len(doc.inline_shapes))
    bundle.result["quality_report"]["excluded"].append("headers_footers_comments_and_revision_markup")


def _parse_pptx(bundle: _Bundle, data: bytes):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(io.BytesIO(data))
    if len(presentation.slides) > MAX_PAGES:
        raise DocumentError("page_limit")

    def shapes(items, slide_index: int, parent_path: str):
        for shape in items:
            location = {"selector_type": "slide_shape", "slide": slide_index + 1, "shape_id": shape.shape_id,
                        "object_path": f"{parent_path}/shape/{shape.shape_id}", "granularity": "shape",
                        "dimensions": [presentation.slide_width, presentation.slide_height],
                        "rects": [[shape.left, shape.top, shape.left + shape.width, shape.top + shape.height]],
                        "coordinate_system": "office_emu_top_left"}
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                shapes(shape.shapes, slide_index, location["object_path"])
            elif shape.has_table:
                grid = [[cell.text for cell in row.cells] for row in shape.table.rows]
                locations = [[{**location, "object_path": f"{location['object_path']}/row/{r}/cell/{c}",
                               "cell_range": f"R{r+1}C{c+1}", "granularity": "cell",
                               "coordinate_quality": "shape_only"} for c in range(len(row))] for r, row in enumerate(grid)]
                merged = [{"row": r+1, "column": c+1, "row_span": cell.span_height, "column_span": cell.span_width}
                          for r, row in enumerate(shape.table.rows) for c, cell in enumerate(row.cells) if cell.is_merge_origin]
                bundle.table(grid, {**location, "granularity": "table"}, locations, merged_ranges=merged)
            elif shape.has_text_frame and shape.text.strip():
                bundle.element(shape.text, location)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                bundle.missing("office_images_not_recognized", slide=slide_index+1, shape_id=shape.shape_id)
            elif getattr(shape, "has_chart", False):
                bundle.missing("office_chart_not_parsed", slide=slide_index+1, shape_id=shape.shape_id)

    for index, slide in enumerate(presentation.slides):
        if slide._element.get("show") == "0" and not bundle.policy.get("include_hidden", False):
            bundle.result["quality_report"]["excluded"].append(f"hidden_slide/{index+1}")
            continue
        shapes(slide.shapes, index, f"slide/{index+1}")
        if slide.has_notes_slide:
            if bundle.policy.get("include_notes", False):
                frame = slide.notes_slide.notes_text_frame
                if frame and frame.text.strip():
                    bundle.element(frame.text, {"selector_type": "object_path", "slide": index+1,
                                                "object_path": f"slide/{index+1}/notes"})
            else:
                bundle.result["quality_report"]["excluded"].append(f"slide_notes/{index+1}")


def _parse_xlsx(bundle: _Bundle, data: bytes):
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string

    workbook = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    include_hidden = bundle.policy.get("include_hidden", False)
    try:
        for sheet in workbook:
            if sheet.sheet_state != "visible" and not include_hidden:
                bundle.result["quality_report"]["excluded"].append(f"hidden_sheet/{sheet.title}")
                continue
            if sheet.max_row * sheet.max_column > MAX_CELLS:
                raise DocumentError("cell_limit")
            rows = [r for r in range(1, sheet.max_row+1) if include_hidden or not sheet.row_dimensions[r].hidden]
            hidden_ranges = [(dimension.min or column_index_from_string(dimension.index),
                              dimension.max or column_index_from_string(dimension.index))
                             for dimension in sheet.column_dimensions.values() if dimension.hidden]
            cols = [c for c in range(1, sheet.max_column+1)
                    if include_hidden or not any(start <= c <= end for start, end in hidden_ranges)]
            grid, locations, details = [], [], {}
            for ri, row in enumerate(rows):
                grid.append([])
                locations.append([])
                for ci, col in enumerate(cols):
                    cell = sheet.cell(row, col)
                    value = _json_value(cell.value)
                    grid[-1].append(value)
                    locations[-1].append({"selector_type": "cell", "sheet": sheet.title,
                                          "cell_range": cell.coordinate, "object_path": f"sheet/{sheet.title}/{cell.coordinate}",
                                          "granularity": "cell"})
                    metadata = {"row": row, "column": col, "number_format": cell.number_format, "data_type": cell.data_type}
                    if cell.data_type == "f":
                        stored = _json_value(cached[sheet.title].cell(row, col).value)
                        metadata.update(formula=value, parsed_value=stored, cached_value=stored,
                                        value_status="cached" if stored is not None else "cached_missing")
                        if stored is None:
                            bundle.warning("formula_cache_missing", sheet=sheet.title, cell=cell.coordinate)
                    details[(ri, ci)] = metadata
            if any(any(value is not None for value in row) for row in grid):
                bundle.table(grid, {"selector_type": "cell_range", "sheet": sheet.title,
                                    "cell_range": sheet.calculate_dimension(), "object_path": f"sheet/{sheet.title}",
                                    "granularity": "table"}, locations, details,
                             merged_ranges=[str(item) for item in sheet.merged_cells.ranges])
            if sheet._images or sheet._charts:
                bundle.missing("office_images_or_charts_not_parsed", sheet=sheet.title)
    finally:
        workbook.close()
        cached.close()


def _parse_image(bundle: _Bundle, data: bytes):
    from PIL import Image

    with Image.open(io.BytesIO(data)) as image:
        image.verify()
    with Image.open(io.BytesIO(data)) as image:
        expected = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".tif": "TIFF", ".tiff": "TIFF"}[bundle.ext]
        if image.format != expected:
            raise DocumentError("file_type_mismatch")
        if image.width * image.height > 50_000_000:
            raise DocumentError("pixel_limit")
        frames = getattr(image, "n_frames", 1)
        if frames > MAX_PAGES:
            raise DocumentError("page_limit")
        bundle.result["quality_report"]["image"] = {
            "dimensions": [image.width, image.height], "frames": frames,
            "exif_orientation": image.getexif().get(274, 1),
        }
    bundle.ocr_needed({"image_frames": frames}, False)


PARSERS = {
    ".txt": _parse_text, ".md": _parse_text, ".csv": _parse_csv,
    ".html": _parse_html, ".htm": _parse_html, ".pdf": _parse_pdf,
    ".docx": _parse_docx, ".pptx": _parse_pptx, ".xlsx": _parse_xlsx,
    ".png": _parse_image, ".jpg": _parse_image, ".jpeg": _parse_image,
    ".tif": _parse_image, ".tiff": _parse_image,
}


def parse_native_document(data: bytes, filename: str, policy: dict | None = None) -> dict:
    """Return a complete bundle or an explicit FAILED/BLOCKED/PARTIAL result."""
    if not isinstance(data, bytes) or not isinstance(filename, str):
        raise TypeError("data must be bytes and filename must be a string")
    bundle = _Bundle(data, filename, policy or {})
    try:
        if not data or len(data) > max_upload_bytes():
            raise DocumentError("empty_or_oversized_input")
        if bundle.ext not in PARSERS:
            raise DocumentError("unsupported_format")
        if not bundle.policy.get("native_parse", True):
            bundle.result["status"] = "BLOCKED"
            bundle.warning("native_parse_blocked_by_policy")
            return bundle.result
        if bundle.ext in (".docx", ".pptx", ".xlsx"):
            _validate_zip(data, bundle.ext)
        PARSERS[bundle.ext](bundle, data)
        if not bundle.result["elements"]:
            if bundle.result["quality_report"]["missing_capabilities"]:
                bundle.result["status"] = "PARTIAL" if bundle.policy.get("ocr", False) else "BLOCKED"
            else:
                bundle.result["status"] = "FAILED"
                bundle.warning("no_extractable_content")
    except Exception as exc:
        bundle.result["status"] = "FAILED"
        code = str(exc) if isinstance(exc, DocumentError) else f"invalid_document_{type(exc).__name__}"
        bundle.warning(code)
        bundle.result["elements"] = []
        bundle.result["evidence_spans"] = []
        bundle.result["anchors"] = []
    return bundle.result


def parse_document(data: bytes, filename: str, policy: dict | None = None) -> dict:
    """Product entrypoint: native parsing followed by the configured layout route."""
    from .parsing_backends import routed_document

    native = parse_native_document(data, filename, policy)
    return routed_document(data, filename, policy or {}, native)


def render_pdf_preview(data: bytes, anchor: dict) -> bytes:
    """Render the immutable source page with its native-coordinate highlight."""
    import pdfplumber

    if hashlib.sha256(data).hexdigest() != anchor.get("source_version_id"):
        raise DocumentError("anchor_source_version_mismatch")
    if anchor.get("selector_type") != "page_region" or anchor.get("coordinate_system") != "pdfplumber_top_left_points":
        raise DocumentError("unsupported_anchor_coordinates")
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        index = anchor.get("page_index")
        if not isinstance(index, int) or index < 0 or index >= len(pdf.pages):
            raise DocumentError("invalid_page_index")
        page = pdf.pages[index]
        image = page.to_image(resolution=110, antialias=True)
        image.draw_rects(anchor.get("rects", []), fill=(255, 210, 0, 60), stroke=(220, 70, 0, 255), stroke_width=2)
        stream = io.BytesIO()
        image.save(stream, format="PNG", quantize=False)
        return stream.getvalue()
